#!/usr/bin/env python3
"""
Generate LeKise CRM Pipeline 2026 - Updated with PL936-PL945
Removes duplicate PL857, adds 9 new June 2026 quotations.
"""

import openpyxl
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              numbers as num_fmt)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, FormulaRule
from openpyxl.styles.differential import DifferentialStyle
from collections import defaultdict
import os

FONT_NAME = "Tahoma"
SRC = "/root/.claude/uploads/8d54d196-3ee6-55bc-9192-3827b3ea07ab/1c393c60-11.06.2569LeKise_CRM_Pipeline_2026_FULL.xlsx"
OUT = "/home/user/peetmaekr/LeKise_CRM_Pipeline_2026_FULL.xlsx"

# ─── Color palette ────────────────────────────────────────────────────────────
C = {
    "header_dark":  "1A237E",
    "header_blue":  "283593",
    "header_teal":  "00695C",
    "header_guard": "4E342E",
    "header_prod":  "37474F",
    "header_fu":    "1B5E20",
    "row_hot":      "FFCCBC",
    "row_high":     "FFF9C4",
    "row_med":      "E8F5E9",
    "row_low":      "FAFAFA",
    "accent_hot":   "D32F2F",
    "accent_high":  "F57F17",
    "accent_med":   "388E3C",
    "accent_low":   "757575",
    "light_blue":   "E3F2FD",
    "light_green":  "E8F5E9",
    "light_amber":  "FFF8E1",
    "white":        "FFFFFF",
    "gray_header":  "ECEFF1",
    "guard_bg":     "FBE9E7",
    "guard_header": "BF360C",
}

def ft(bold=False, size=10, color="000000", name=FONT_NAME, italic=False):
    return Font(name=name, bold=bold, size=size, color=color, italic=italic)

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def al(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def border(style="thin", color="BDBDBD"):
    s = Side(style=style, color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def thick_border():
    tk = Side(style="medium", color="546E7A")
    th = Side(style="thin", color="BDBDBD")
    return Border(left=tk, right=tk, top=th, bottom=th)

def num(n, dec=2):
    return round(float(n or 0), dec)

# ─── Load existing data ───────────────────────────────────────────────────────
print("Loading existing data...")
wb_src = openpyxl.load_workbook(SRC)
ws_src = wb_src["📋 CRM Table"]

existing_rows = []
for row in ws_src.iter_rows(min_row=5, max_row=ws_src.max_row, values_only=True):
    if row[0] and str(row[0]).startswith("PL"):
        existing_rows.append(list(row))

print(f"  Loaded {len(existing_rows)} existing records")

# ─── New records from PDFs PL936-PL945 ───────────────────────────────────────
# Cols: PL, Month, Customer, Type, QuoteNo, Date, ProductGroup, ProductDetail,
#        BeforeVAT, VAT7, GrandTotal, StatusAuto, Priority, Count, BusinessGroup,
#        StatusReal, Notes, FollowDate, Action
new_records = [
    ("PL936","2026-06","คุณนพกร","B2C บุคคล","LKT. 2026-06 -PL936","12 มิถุนายน 2026",
     "Guard Rail / การ์ดเรล","การ์ดเรล Ø48.3×2.5t สีขาว-แดง",
     299651.52,20975.61,320627.13,"📋 ติดตาม","กลาง",1,"—",None,None,None,None),

    ("PL938","2026-06","เทศบาลเมืองตราด","B2G อปท.","LKT. 2026-06 -PL938","16 มิถุนายน 2026",
     "High Mast","High Mast 15M+Floodlight 400W+HDPE Pipe+Cable",
     1241571.80,86910.03,1328481.83,"📞 ติดตาม-High","สูง",1,"—",None,None,None,None),

    ("PL939","2026-06","องค์การบริหารส่วนตำบลตาขัน","B2G อปท.","LKT. 2026-06 -PL939","12 มิถุนายน 2026",
     "Guard Rail / การ์ดเรล","การ์ดเรล Ø60.3×3.2t สีเทา",
     845503.41,59185.24,904688.65,"📋 ติดตาม","กลาง",1,"—",None,None,None,None),

    ("PL940","2026-06","บริษัท โพรเซสเอ็นจิเนียรแอนดคอนซัลแทนท จำกัด","B2B บริษัท","LKT. 2026-06 -PL940","15 มิถุนายน 2026",
     "High Mast","High Mast 20M+Floodlight 400W (REV.ใหม่ แทน PL857)",
     1009576.00,70670.32,1080246.32,"📞 ติดตาม-High","สูง",1,"—",None,None,None,None),

    ("PL941","2026-06","นิติบุคคลหมู่บ้านจัดสรรเดอะลากูน","B2B อื่นๆ","LKT. 2026-06 -PL941","15 มิถุนายน 2026",
     "เสาไฟถนน","Post Top 6M+Lamp",
     60799.20,4255.94,65055.14,"⚪ รอผล","ต่ำ",1,"—",None,None,None,None),

    ("PL942","2026-06","สำนักงานเทศบาลนครมาบตาพุด","B2G อปท.","LKT. 2026-06 -PL942","15 มิถุนายน 2026",
     "LED Streetlight","LED Streetlight 120W×19 ชุด",
     96824.00,6777.68,103601.68,"⚪ รอผล","ต่ำ",1,"—",None,None,None,None),

    ("PL943","2026-06","บมจ.เมืองไทยประกันชีวิต","B2B บริษัท","LKT. 2026-06 -PL943","16 มิถุนายน 2026",
     "High Mast","High Mast 15M+Motor (4 ต้น)+ค่าขนส่ง",
     610599.00,42741.93,653340.93,"📋 ติดตาม","กลาง",1,"—",None,None,None,None),

    ("PL944","2026-06","บริษัท ทีเจริญ บิลด จำกัด","B2B บริษัท","LKT. 2026-06 -PL944","16 มิถุนายน 2026",
     "Guard Rail / การ์ดเรล","การ์ดเรล Ø48.3×2.5t (small lot)+ค่าส่ง",
     15765.28,1103.57,16868.85,"⚪ รอผล","ต่ำ",1,"—",None,None,None,None),

    ("PL945","2026-06","บริษัท แอสเซนท คอรปอเรชั่น จำกัด","B2B บริษัท","LKT. 2026-06 -PL945","16 มิถุนายน 2026",
     "High Mast","High Mast 15M วงแหวน 6 โคม",
     221186.00,15483.02,236669.02,"⚪ รอผล","ต่ำ",1,"—",None,None,None,None),
]

# ─── Duplicate removal ────────────────────────────────────────────────────────
REMOVE_PLs = {"PL857"}  # Replaced by PL940 (same customer, same product, diff 3.6%)
removed_info = [r for r in existing_rows if r[0] in REMOVE_PLs]

all_data = [r for r in existing_rows if r[0] not in REMOVE_PLs]
all_data.extend([list(r) for r in new_records])
all_data.sort(key=lambda r: int(str(r[0]).replace("PL", "")))

total_records = len(all_data)
print(f"  Removed {len(REMOVE_PLs)} duplicate(s): {REMOVE_PLs}")
print(f"  Added {len(new_records)} new records")
print(f"  Total records: {total_records}")

# ─── Summary statistics ───────────────────────────────────────────────────────
grand_total_all    = sum(r[10] or 0 for r in all_data)
before_vat_all     = sum(r[8]  or 0 for r in all_data)
hot_total          = sum(r[10] or 0 for r in all_data if "ร้อน"   in str(r[11]))
high_total         = sum(r[10] or 0 for r in all_data if "High"   in str(r[11]))
follow_total       = sum(r[10] or 0 for r in all_data if "📋"     in str(r[11]))
wait_total         = sum(r[10] or 0 for r in all_data if "⚪"     in str(r[11]))
unique_customers   = len({r[2] for r in all_data})

# By month
by_month = defaultdict(lambda: {"count": 0, "total": 0.0})
for r in all_data:
    by_month[r[1]]["count"] += 1
    by_month[r[1]]["total"] += r[10] or 0

# By product group
by_product = defaultdict(lambda: {"count": 0, "total": 0.0})
for r in all_data:
    by_product[r[6]]["count"] += 1
    by_product[r[6]]["total"] += r[10] or 0

# By business group
by_bg = defaultdict(lambda: {"count": 0, "total": 0.0})
for r in all_data:
    bg = r[14] if r[14] and r[14] != "—" else "อื่นๆ / ทั่วไป"
    by_bg[bg]["count"] += 1
    by_bg[bg]["total"] += r[10] or 0

# By customer (for top 20)
by_customer = defaultdict(lambda: {"count": 0, "total": 0.0, "products": set()})
for r in all_data:
    by_customer[r[2]]["count"] += 1
    by_customer[r[2]]["total"] += r[10] or 0
    by_customer[r[2]]["products"].add(r[6])

MONTHS = ["2026-01","2026-02","2026-03","2026-04","2026-05","2026-06"]
PRODUCTS_ORDER = ["High Mast","Guard Rail / การ์ดเรล","เสาไฟถนน","LED Streetlight",
                  "LED Floodlight","LED Highbay","Solar","สายส่งไฟฟ้า",
                  "Hardware / Accessories","Electrical Accessories","อื่นๆ"]

# ─── Build workbook ───────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
wb.remove(wb.active)

def set_col_widths(ws, widths):
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w

def write_title_row(ws, title, row=1, merge_range=None, height=22):
    if merge_range is None:
        merge_range = f"A{row}:N{row}"
    ws.merge_cells(merge_range)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = ft(bold=True, size=13, color="FFFFFF")
    cell.fill = fill(C["header_dark"])
    cell.alignment = al("center")
    ws.row_dimensions[row].height = height

def header_row(ws, headers, row, colors=None, font_color="FFFFFF"):
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = ft(bold=True, size=9, color=font_color)
        c.fill = fill(colors[col-1] if colors else C["header_blue"])
        c.alignment = al("center", wrap=True)
        c.border = border("thin", "7B8D9B")
    ws.row_dimensions[row].height = 30

# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 1: 📊 Summary Dashboard
# ═══════════════════════════════════════════════════════════════════════════════
print("Building Summary Dashboard...")
ws1 = wb.create_sheet("📊 Summary Dashboard")
set_col_widths(ws1, [22,18,18,18,18,18,18,18,18,16,16,16,16,16])

write_title_row(ws1, f"  LeKise Trading  —  CRM & Pipeline Analysis 2026  |  Pole Division  |  อัปเดต 16 มิ.ย. 2569", 1, "A1:N1")

ws1.row_dimensions[2].height = 6

# KPI row 3 - labels
kpi_labels = [
    ("Pipeline รวม (incl VAT)", C["header_dark"]),
    ("ก่อน VAT รวม", C["header_blue"]),
    ("🔥 Hot Deals (≥3M)", "B71C1C"),
    ("📞 Follow-High (1-3M)", "E65100"),
    ("จำนวนใบเสนอราคา", "37474F"),
    ("ลูกค้าไม่ซ้ำ", "1A237E"),
    ("📋 ติดตาม (0.3-1M)", "2E7D32"),
    ("⚪ รอผล (<0.3M)", "546E7A"),
]
for i, (label, bg) in enumerate(kpi_labels):
    col = i * 2 + 1
    if col + 1 <= 16:
        ws1.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col+1)
    c = ws1.cell(row=3, column=col, value=f"  {label}")
    c.font = ft(bold=True, size=9, color="FFFFFF")
    c.fill = fill(bg)
    c.alignment = al("left")
    ws1.row_dimensions[3].height = 20

# KPI row 4 - values
kpi_vals = [
    (f"฿ {grand_total_all/1e6:.1f}M", C["header_dark"]),
    (f"฿ {before_vat_all/1e6:.1f}M", C["header_blue"]),
    (f"฿ {hot_total/1e6:.1f}M", "B71C1C"),
    (f"฿ {high_total/1e6:.1f}M", "E65100"),
    (f"{total_records} ใบ", "37474F"),
    (f"{unique_customers} ราย", "1A237E"),
    (f"฿ {follow_total/1e6:.1f}M", "2E7D32"),
    (f"฿ {wait_total/1e6:.1f}M", "546E7A"),
]
for i, (val, bg) in enumerate(kpi_vals):
    col = i * 2 + 1
    if col + 1 <= 16:
        ws1.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col+1)
    c = ws1.cell(row=4, column=col, value=val)
    c.font = ft(bold=True, size=14, color="FFFFFF")
    c.fill = fill(bg)
    c.alignment = al("center")
    ws1.row_dimensions[4].height = 28

ws1.row_dimensions[5].height = 8

# Monthly Pipeline table (rows 6-15)
ws1.merge_cells("A6:B6")
t = ws1.cell(row=6, column=1, value="📅 Pipeline รายเดือน (incl VAT)")
t.font = ft(bold=True, size=10, color="FFFFFF")
t.fill = fill(C["header_blue"])
t.alignment = al("center")

month_headers = ["เดือน","ใบเสนอ","ก่อน VAT","VAT 7%","รวม incl VAT","% Pipeline","สถานะ"]
for col, h in enumerate(month_headers, start=1):
    c = ws1.cell(row=7, column=col, value=h)
    c.font = ft(bold=True, size=9, color="FFFFFF")
    c.fill = fill(C["header_teal"])
    c.alignment = al("center")
    c.border = border()
ws1.row_dimensions[7].height = 22

row_idx = 8
for m in MONTHS:
    d = by_month.get(m, {"count": 0, "total": 0.0})
    cnt = d["count"]
    tot = d["total"]
    bv = tot / 1.07
    vat = tot - bv
    pct = tot / grand_total_all * 100 if grand_total_all else 0
    bg = C["light_blue"] if row_idx % 2 == 0 else C["white"]
    vals = [m, cnt, round(bv, 2), round(vat, 2), round(tot, 2), f"{pct:.1f}%",
            "🔥" if tot >= 3e6 else "📞" if tot >= 1e6 else "📋" if tot >= 3e5 else "⚪"]
    for col, v in enumerate(vals, start=1):
        c = ws1.cell(row=row_idx, column=col, value=v)
        c.font = ft(size=9)
        c.fill = fill(bg)
        c.alignment = al("center")
        c.border = border()
        if col in (3, 4, 5) and isinstance(v, float):
            c.number_format = "#,##0.00"
    row_idx += 1

# Total row
ws1.merge_cells(f"A{row_idx}:B{row_idx}")
c = ws1.cell(row=row_idx, column=1, value="รวมทั้งหมด")
c.font = ft(bold=True, size=9)
c.fill = fill(C["header_dark"])
c.font = Font(name=FONT_NAME, bold=True, size=9, color="FFFFFF")
c.alignment = al("center")
c.border = border()
for col, val in zip([3, 4, 5, 6], [round(before_vat_all, 2), round(grand_total_all - before_vat_all, 2), round(grand_total_all, 2), "100%"]):
    c2 = ws1.cell(row=row_idx, column=col, value=val)
    c2.font = ft(bold=True, size=9, color="FFFFFF")
    c2.fill = fill(C["header_dark"])
    c2.alignment = al("center")
    c2.border = border()
    if col in (3, 4, 5) and isinstance(val, float):
        c2.number_format = "#,##0.00"
ws1.row_dimensions[row_idx].height = 20
row_idx += 2

# Product Group table
ws1.merge_cells(f"A{row_idx}:C{row_idx}")
t = ws1.cell(row=row_idx, column=1, value="📦 Pipeline by กลุ่มสินค้า")
t.font = ft(bold=True, size=10, color="FFFFFF")
t.fill = fill("4A148C")
t.alignment = al("center")
row_idx += 1

ph = ["กลุ่มสินค้า","ใบเสนอ","มูลค่า (incl VAT)","% Pipeline"]
for col, h in enumerate(ph, start=1):
    c = ws1.cell(row=row_idx, column=col, value=h)
    c.font = ft(bold=True, size=9, color="FFFFFF")
    c.fill = fill("6A1B9A")
    c.alignment = al("center")
    c.border = border()
row_idx += 1

for pg in PRODUCTS_ORDER:
    d = by_product.get(pg, {"count": 0, "total": 0.0})
    if d["count"] == 0:
        continue
    pct = d["total"] / grand_total_all * 100 if grand_total_all else 0
    bg = C["white"] if row_idx % 2 == 0 else C["light_blue"]
    for col, v in enumerate([pg, d["count"], round(d["total"], 2), f"{pct:.1f}%"], start=1):
        c = ws1.cell(row=row_idx, column=col, value=v)
        c.font = ft(size=9)
        c.fill = fill(bg)
        c.alignment = al("center" if col > 1 else "left")
        c.border = border()
        if col == 3:
            c.number_format = "#,##0.00"
    row_idx += 1

row_idx += 1

# Top 20 customers
ws1.merge_cells(f"A{row_idx}:E{row_idx}")
t = ws1.cell(row=row_idx, column=1, value="🏆 Top 20 ลูกค้า (Pipeline incl VAT)")
t.font = ft(bold=True, size=10, color="FFFFFF")
t.fill = fill("B71C1C")
t.alignment = al("center")
row_idx += 1

top20_headers = ["#","ชื่อลูกค้า","ใบเสนอ","มูลค่า (฿)","% Pipeline"]
for col, h in enumerate(top20_headers, start=1):
    c = ws1.cell(row=row_idx, column=col, value=h)
    c.font = ft(bold=True, size=9, color="FFFFFF")
    c.fill = fill("C62828")
    c.alignment = al("center")
    c.border = border()
row_idx += 1

top20 = sorted(by_customer.items(), key=lambda x: x[1]["total"], reverse=True)[:20]
for rank, (cust, d) in enumerate(top20, start=1):
    pct = d["total"] / grand_total_all * 100 if grand_total_all else 0
    bg = C["row_hot"] if d["total"] >= 3e6 else C["row_high"] if d["total"] >= 1e6 else C["white"]
    for col, v in enumerate([rank, cust, d["count"], round(d["total"], 2), f"{pct:.1f}%"], start=1):
        c = ws1.cell(row=row_idx, column=col, value=v)
        c.font = ft(size=9)
        c.fill = fill(bg)
        c.alignment = al("center" if col != 2 else "left")
        c.border = border()
        if col == 4:
            c.number_format = "#,##0.00"
    row_idx += 1

ws1.sheet_view.zoomScale = 90
ws1.freeze_panes = "A5"

# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 2: 📋 CRM Table
# ═══════════════════════════════════════════════════════════════════════════════
print("Building CRM Table...")
ws2 = wb.create_sheet("📋 CRM Table")
set_col_widths(ws2, [8,10,26,12,22,20,20,38,14,12,14,18,10,8,14,20,28,14,14])

write_title_row(ws2,
    f"  LeKise CRM  |  ใบเสนอราคา 2026  |  {total_records} รายการ  |  อัปเดต 16 มิ.ย. 2569  |  ตัด Revision ซ้ำ (รวม PL857)",
    row=1, merge_range="A1:S1")
ws2.row_dimensions[2].height = 6

# Summary row 3
kpi_crm = [
    (f"Pipeline: {grand_total_all/1e6:.1f}M ฿", 5),
    (f"Hot≥3M: {hot_total/1e6:.1f}M ฿", 5),
    (f"{total_records} ใบ (ทบทวนแล้ว)", 4),
    (f"EGAT:{by_bg['EGAT']['total']/1e6:.1f}M | DOA:{by_bg['DOA']['total']/1e6:.1f}M | APT:{by_bg['อปท. Supply']['total']/1e6:.1f}M", 5),
]
col = 1
for text, span in kpi_crm:
    ws2.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + span - 1)
    c = ws2.cell(row=3, column=col, value=f"  {text}")
    c.font = ft(bold=True, size=9, color="1A237E")
    c.fill = fill(C["light_blue"])
    c.alignment = al("left")
    col += span
ws2.row_dimensions[3].height = 18

# Header row 4
crm_headers = ["PL No.","เดือน","ชื่อลูกค้า","ประเภท","เลขที่ใบเสนอ","วันที่",
                "กลุ่มสินค้า","สินค้าหลัก","ก่อน VAT (฿)","VAT 7% (฿)","รวม incl VAT (฿)",
                "สถานะ (Auto)","ความสำคัญ","จำนวน","กลุ่มงาน",
                "★ สถานะตามจริง","หมายเหตุ / ผลติดต่อ","วันติดตาม","Action"]
header_row(ws2, crm_headers, 4)
ws2.row_dimensions[4].height = 32

# Data rows
status_color_map = {
    "🔥": C["row_hot"],
    "📞": C["row_high"],
    "📋": C["row_med"],
    "⚪": C["row_low"],
}

for r_idx, row in enumerate(all_data, start=5):
    status = str(row[11] or "")
    bg = C["row_low"]
    for k, v in status_color_map.items():
        if k in status:
            bg = v
            break

    for col, val in enumerate(row, start=1):
        c = ws2.cell(row=r_idx, column=col, value=val)
        c.font = ft(size=9)
        c.fill = fill(bg)
        c.border = border("thin", "CFD8DC")
        if col == 3:
            c.alignment = al("left", wrap=False)
        elif col == 8:
            c.alignment = al("left", wrap=True)
        elif col in (9, 10, 11):
            c.alignment = al("right")
            c.number_format = "#,##0.00"
        else:
            c.alignment = al("center")
    ws2.row_dimensions[r_idx].height = 15

# Grand total row
gt_row = 5 + total_records
ws2.merge_cells(f"A{gt_row}:H{gt_row}")
c = ws2.cell(row=gt_row, column=1, value=f"GRAND TOTAL  |  {total_records} รายการ  |  ทบทวนแล้ว 16 มิ.ย. 2569")
c.font = ft(bold=True, size=9, color="FFFFFF")
c.fill = fill(C["header_dark"])
c.alignment = al("left")

sum_before_vat = sum(r[8] or 0 for r in all_data)
sum_vat        = sum(r[9] or 0 for r in all_data)
sum_total      = sum(r[10] or 0 for r in all_data)
for col, val in [(9, sum_before_vat), (10, sum_vat), (11, sum_total)]:
    c2 = ws2.cell(row=gt_row, column=col, value=round(val, 2))
    c2.font = ft(bold=True, size=9, color="FFFFFF")
    c2.fill = fill(C["header_dark"])
    c2.alignment = al("center")
    c2.number_format = "#,##0.00"
ws2.row_dimensions[gt_row].height = 20

# Data validation for column P (★ สถานะตามจริง)
status_options = '"✅ปิดได้,🔥ร้อน-ปิดเร็ว,📞กำลังเจรจา,💼ส่ง Proposal แล้ว,⏳รอลูกค้าตอบ,❌พัก/ไม่ซื้อ,🔄Revise ราคา,⚪ยังไม่ติดต่อ"'
dv = DataValidation(type="list", formula1=status_options, allow_blank=True)
dv.error = "กรุณาเลือกจาก Dropdown"
dv.prompt = "เลือกสถานะตามจริง"
ws2.add_data_validation(dv)
dv.sqref = f"P5:P{gt_row - 1}"

ws2.freeze_panes = "C5"
ws2.sheet_view.zoomScale = 90

# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 3: 📈 Status Dashboard (COUNTIF/SUMIF formulas referencing CRM Table col P)
# ═══════════════════════════════════════════════════════════════════════════════
print("Building Status Dashboard...")
ws3 = wb.create_sheet("📈 Status Dashboard")
set_col_widths(ws3, [28,14,18,18,16,16,16,16,16,16,16,14,14,14])

write_title_row(ws3,
    "  LeKise  |  STATUS DASHBOARD  |  อัปเดตอัตโนมัติจาก CRM Table  |  กรอกสถานะตามจริงที่ Column P ใน CRM Table",
    1, "A1:N1")
ws3.row_dimensions[2].height = 6

note_row = 3
ws3.merge_cells("A3:N3")
c = ws3.cell(row=note_row, column=1,
             value="  💡  วิธีใช้: ไปที่ Sheet 📋 CRM Table → คอลัมน์ P \"★ สถานะตามจริง\" → เลือก Dropdown  |  Dashboard นี้จะ Update อัตโนมัติ")
c.font = ft(size=9, color="0D47A1", italic=True)
c.fill = fill(C["light_blue"])
c.alignment = al("left")
ws3.row_dimensions[3].height = 18

crm_ref_col = "P"
crm_data_range = f"'📋 CRM Table'!{crm_ref_col}5:{crm_ref_col}{gt_row - 1}"
crm_total_range = f"'📋 CRM Table'!K5:K{gt_row - 1}"

status_real_list = [
    ("✅ปิดได้",        "1B5E20", "27AE60"),
    ("🔥ร้อน-ปิดเร็ว",  "B71C1C", "E53935"),
    ("📞กำลังเจรจา",    "E65100", "FB8C00"),
    ("💼ส่ง Proposal แล้ว", "1A237E", "1E88E5"),
    ("⏳รอลูกค้าตอบ",   "4A148C", "8E24AA"),
    ("❌พัก/ไม่ซื้อ",   "37474F", "546E7A"),
    ("🔄Revise ราคา",   "33691E", "7CB342"),
    ("⚪ยังไม่ติดต่อ",   "455A64", "78909C"),
]

row5_labels = ["สถานะตามจริง", "จำนวนใบ", "มูลค่ารวม (฿)", "% มูลค่า", "% ใบ", "เฉลี่ย/ใบ"]
for col, h in enumerate(row5_labels, start=1):
    c = ws3.cell(row=5, column=col, value=h)
    c.font = ft(bold=True, size=9, color="FFFFFF")
    c.fill = fill(C["header_dark"])
    c.alignment = al("center")
    c.border = border()
ws3.row_dimensions[5].height = 24

for r_i, (status_name, dark, light) in enumerate(status_real_list, start=6):
    cnt_formula  = f'=COUNTIF({crm_data_range},"{status_name}")'
    sum_formula  = f'=SUMIF({crm_data_range},"{status_name}",{crm_total_range})'
    pct_val_formula = f'=IF(SUM({crm_total_range})=0,0,{get_column_letter(3)}{r_i}/SUM({crm_total_range}))'
    pct_cnt_formula = f'=IF(COUNTA({crm_data_range})=0,0,{get_column_letter(2)}{r_i}/COUNTA({crm_data_range}))'
    avg_formula  = f'=IF({get_column_letter(2)}{r_i}=0,0,{get_column_letter(3)}{r_i}/{get_column_letter(2)}{r_i})'
    row_data = [status_name, cnt_formula, sum_formula, pct_val_formula, pct_cnt_formula, avg_formula]
    bg = C["light_green"] if r_i % 2 == 0 else C["white"]
    for col, val in enumerate(row_data, start=1):
        c = ws3.cell(row=r_i, column=col, value=val)
        c.font = ft(size=9)
        c.fill = fill(bg)
        c.border = border()
        if col == 1:
            c.font = ft(bold=True, size=9, color=dark)
            c.fill = fill(bg)
        elif col == 3:
            c.number_format = "#,##0.00"
        elif col in (4, 5):
            c.number_format = "0.0%"
        elif col == 6:
            c.number_format = "#,##0.00"
        c.alignment = al("center" if col > 1 else "left")
    ws3.row_dimensions[r_i].height = 18

# Total row
total_row3 = 6 + len(status_real_list)
ws3.cell(row=total_row3, column=1, value="รวมทั้งหมด (ที่กรอกสถานะ)").font = ft(bold=True, size=9)
ws3.cell(row=total_row3, column=2, value=f'=COUNTA({crm_data_range})').number_format = "#,##0"
ws3.cell(row=total_row3, column=3, value=f'=SUM({crm_total_range})').number_format = "#,##0.00"
for col in range(1, 7):
    c = ws3.cell(row=total_row3, column=col)
    c.font = ft(bold=True, size=9, color="FFFFFF")
    c.fill = fill(C["header_dark"])
    c.alignment = al("center" if col > 1 else "left")
    c.border = border()

# Auto status breakdown (from col L)
r_start = total_row3 + 2
ws3.merge_cells(f"A{r_start}:F{r_start}")
c = ws3.cell(row=r_start, column=1, value="📊 สถานะ Auto (จาก Pipeline Value) — ข้อมูลอ้างอิงจาก CRM Table คอลัมน์ L")
c.font = ft(bold=True, size=10, color="FFFFFF")
c.fill = fill(C["header_teal"])
c.alignment = al("center")
ws3.row_dimensions[r_start].height = 20

auto_status_list = [
    ("🔥 ร้อน", "≥ 3,000,000 ฿", C["row_hot"]),
    ("📞 ติดตาม-High", "1,000,000 – 2,999,999 ฿", C["row_high"]),
    ("📋 ติดตาม", "300,000 – 999,999 ฿", C["row_med"]),
    ("⚪ รอผล", "< 300,000 ฿", C["row_low"]),
]
crm_l_range = f"'📋 CRM Table'!L5:L{gt_row - 1}"
for r_i, (st, rng, bg) in enumerate(auto_status_list, start=r_start + 1):
    cnt_f = f'=COUNTIF({crm_l_range},"*{st.split()[0]}*")'
    sum_f = f'=SUMIF({crm_l_range},"*{st.split()[0]}*",{crm_total_range})'
    for col, val in enumerate([st, rng, cnt_f, sum_f, "", ""], start=1):
        c = ws3.cell(row=r_i, column=col, value=val)
        c.font = ft(size=9)
        c.fill = fill(bg)
        c.alignment = al("center" if col > 2 else "left")
        c.border = border()
        if col == 4:
            c.number_format = "#,##0.00"
    ws3.row_dimensions[r_i].height = 16

ws3.freeze_panes = "A6"
ws3.sheet_view.zoomScale = 90

# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 4: 👥 Customer CRM
# ═══════════════════════════════════════════════════════════════════════════════
print("Building Customer CRM...")
ws4 = wb.create_sheet("👥 Customer CRM")
set_col_widths(ws4, [5,30,14,8,16,8,22,10,12,14,16,22])

write_title_row(ws4,
    f"  LeKise  |  Customer CRM 2026  |  {unique_customers} ลูกค้า  |  16 มิ.ย. 2569",
    1, "A1:L1")
ws4.row_dimensions[2].height = 6

cust_headers = ["#","ชื่อลูกค้า","ประเภท","Deal","มูลค่า (฿)","%","กลุ่มสินค้า","เดือน","PL","ระดับ CRM","กลุ่มงาน","Action"]
header_row(ws4, cust_headers, 3)
ws4.row_dimensions[3].height = 28

# Build customer summary
cust_detail = defaultdict(lambda: {"count":0,"total":0.0,"products":set(),"months":set(),"pl_list":[],"type":"","bg":"—"})
for r in all_data:
    cust = r[2]
    cust_detail[cust]["count"] += 1
    cust_detail[cust]["total"] += r[10] or 0
    cust_detail[cust]["products"].add(r[6])
    cust_detail[cust]["months"].add(r[1])
    cust_detail[cust]["pl_list"].append(r[0])
    cust_detail[cust]["type"] = r[3]
    if r[14] and r[14] != "—":
        cust_detail[cust]["bg"] = r[14]

sorted_custs = sorted(cust_detail.items(), key=lambda x: x[1]["total"], reverse=True)
for i, (cust, d) in enumerate(sorted_custs, start=1):
    tot = d["total"]
    pct = tot / grand_total_all * 100 if grand_total_all else 0
    crm_level = "A - VIP" if tot >= 3e6 else "B - สำคัญ" if tot >= 1e6 else "C - ทั่วไป" if tot >= 3e5 else "D - รอผล"
    action = "📞 ติดตามด่วน" if tot >= 3e6 else "📋 นัดหมาย" if tot >= 1e6 else "📧 ส่ง Email" if tot >= 3e5 else "⚪ รอสอบถาม"
    bg = C["row_hot"] if tot >= 3e6 else C["row_high"] if tot >= 1e6 else C["row_med"] if tot >= 3e5 else C["row_low"]
    row_i = i + 3
    vals = [i, cust, d["type"], d["count"], round(tot, 2), f"{pct:.1f}%",
            "/".join(list(d["products"])[:2]), ", ".join(sorted(d["months"])[-2:]),
            f'PL{d["pl_list"][-1].replace("PL","") if d["pl_list"] else ""}',
            crm_level, d["bg"], action]
    for col, val in enumerate(vals, start=1):
        c = ws4.cell(row=row_i, column=col, value=val)
        c.font = ft(size=9)
        c.fill = fill(bg)
        c.border = border("thin", "CFD8DC")
        c.alignment = al("left" if col in (2, 7, 12) else "center")
        if col == 5:
            c.number_format = "#,##0.00"
    ws4.row_dimensions[row_i].height = 14

ws4.freeze_panes = "B4"
ws4.sheet_view.zoomScale = 90

# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 5: 🏢 Business Groups
# ═══════════════════════════════════════════════════════════════════════════════
print("Building Business Groups...")
ws5 = wb.create_sheet("🏢 Business Groups")
set_col_widths(ws5, [8,26,14,20,16,14,16])

write_title_row(ws5,
    "  LeKise  |  กลุ่มงาน EGAT / DOA / ZETA-iOT / อปท. Supply  |  16 มิ.ย. 2569",
    1, "A1:G1")
ws5.row_dimensions[2].height = 6

bg_groups = [
    ("EGAT",        "🏭 กลุ่มงาน EGAT (การไฟฟ้าฝ่ายผลิต) — รวม บ.โกลว์ พาวเวอร์ เทคโนโลยี จำกัด", "1565C0", "E3F2FD"),
    ("DOA",         "🌾 กลุ่มงาน DOA (กรมวิชาการเกษตร) — ควินตัส / เดลต้า พลัส / เจมินี เทค",       "2E7D32", "E8F5E9"),
    ("ZETA-iOT",    "📡 กลุ่มงาน ZETA-iOT — ซีซีทีวี / ZETA-9 Cloud / คาแนล วัน",                   "6A1B9A", "F3E5F5"),
    ("อปท. Supply", "🏛️ กลุ่มงาน อปท. Supply — เค บี เอ็ม เทคโนโลยี จำกัด",                         "E65100", "FFF3E0"),
]

row_ptr = 3
bg_headers = ["#","ชื่อลูกค้า","เดือน","สินค้าหลัก","PL No.","ก่อน VAT","รวม incl VAT"]

for bg_key, bg_title, hdr_color, row_bg in bg_groups:
    bg_rows = [r for r in all_data if r[14] == bg_key]
    bg_sum = sum(r[10] or 0 for r in bg_rows)

    ws5.merge_cells(f"A{row_ptr}:G{row_ptr}")
    c = ws5.cell(row=row_ptr, column=1, value=f"{bg_title}  |  {len(bg_rows)} รายการ  |  {bg_sum/1e6:.2f}M ฿")
    c.font = ft(bold=True, size=10, color="FFFFFF")
    c.fill = fill(hdr_color)
    c.alignment = al("left")
    ws5.row_dimensions[row_ptr].height = 18
    row_ptr += 1

    for col, h in enumerate(bg_headers, start=1):
        c = ws5.cell(row=row_ptr, column=col, value=h)
        c.font = ft(bold=True, size=9, color="FFFFFF")
        c.fill = fill(hdr_color)
        c.alignment = al("center")
        c.border = border()
    ws5.row_dimensions[row_ptr].height = 22
    row_ptr += 1

    for i, r in enumerate(bg_rows, start=1):
        bg_cell = row_bg if i % 2 == 0 else C["white"]
        for col, val in enumerate([i, r[2], r[1], r[7][:35] if r[7] else "", r[0], round(r[8] or 0, 2), round(r[10] or 0, 2)], start=1):
            c = ws5.cell(row=row_ptr, column=col, value=val)
            c.font = ft(size=9)
            c.fill = fill(bg_cell)
            c.border = border("thin", "CFD8DC")
            c.alignment = al("left" if col in (2, 4) else "center")
            if col in (6, 7):
                c.number_format = "#,##0.00"
        ws5.row_dimensions[row_ptr].height = 14
        row_ptr += 1

    # Sub total
    ws5.merge_cells(f"A{row_ptr}:E{row_ptr}")
    c = ws5.cell(row=row_ptr, column=1, value=f"รวม {bg_key}")
    c.font = ft(bold=True, size=9, color="FFFFFF")
    c.fill = fill(hdr_color)
    c.alignment = al("center")
    for col, val in [(6, round(sum(r[8] or 0 for r in bg_rows), 2)),
                     (7, round(bg_sum, 2))]:
        c2 = ws5.cell(row=row_ptr, column=col, value=val)
        c2.font = ft(bold=True, size=9, color="FFFFFF")
        c2.fill = fill(hdr_color)
        c2.number_format = "#,##0.00"
        c2.alignment = al("center")
    ws5.row_dimensions[row_ptr].height = 16
    row_ptr += 2

ws5.freeze_panes = "A3"
ws5.sheet_view.zoomScale = 90

# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 6: 🛡️ Guard Rail
# ═══════════════════════════════════════════════════════════════════════════════
print("Building Guard Rail sheet...")
ws6 = wb.create_sheet("🛡️ Guard Rail")
set_col_widths(ws6, [8,26,10,10,35,14,12,14,14,14])

gr_rows = [r for r in all_data if "Guard Rail" in str(r[6])]
gr_total = sum(r[10] or 0 for r in gr_rows)

write_title_row(ws6,
    f"  LeKise  |  Guard Rail / การ์ดเรล  |  {len(gr_rows)} รายการ  |  {gr_total/1e6:.2f}M ฿  |  16 มิ.ย. 2569",
    1, "A1:J1")
ws6.row_dimensions[2].height = 6

ws6.merge_cells("A3:J3")
c = ws6.cell(row=3, column=1, value=f"Pipeline Guard Rail รวม: {gr_total/1e6:.2f}M ฿  |  {len(gr_rows)} รายการ")
c.font = ft(bold=True, size=10, color=C["guard_header"])
c.fill = fill(C["guard_bg"])
c.alignment = al("center")
ws6.row_dimensions[3].height = 18

gr_headers = ["#","ชื่อลูกค้า","เดือน","PL No.","สินค้า","ก่อน VAT (฿)","VAT 7%","รวม incl VAT","สถานะ (Auto)","กลุ่มงาน"]
for col, h in enumerate(gr_headers, start=1):
    c = ws6.cell(row=4, column=col, value=h)
    c.font = ft(bold=True, size=9, color="FFFFFF")
    c.fill = fill(C["guard_header"])
    c.alignment = al("center", wrap=True)
    c.border = border()
ws6.row_dimensions[4].height = 28

for i, r in enumerate(sorted(gr_rows, key=lambda x: x[10] or 0, reverse=True), start=1):
    row_i = i + 4
    bg = C["guard_bg"] if i % 2 == 0 else C["white"]
    status = str(r[11] or "")
    if "🔥" in status: bg = C["row_hot"]
    elif "📞" in status: bg = C["row_high"]
    vals = [i, r[2], r[1], r[0], r[7], round(r[8] or 0, 2), round(r[9] or 0, 2),
            round(r[10] or 0, 2), r[11], r[14] or "—"]
    for col, val in enumerate(vals, start=1):
        c = ws6.cell(row=row_i, column=col, value=val)
        c.font = ft(size=9)
        c.fill = fill(bg)
        c.border = border("thin", "CFD8DC")
        c.alignment = al("left" if col in (2, 5) else "center")
        if col in (6, 7, 8):
            c.number_format = "#,##0.00"
    ws6.row_dimensions[row_i].height = 14

# Total row
total_row6 = len(gr_rows) + 5
ws6.merge_cells(f"A{total_row6}:E{total_row6}")
c = ws6.cell(row=total_row6, column=1, value=f"รวม Guard Rail / การ์ดเรล — {len(gr_rows)} รายการ")
c.font = ft(bold=True, size=9, color="FFFFFF")
c.fill = fill(C["guard_header"])
c.alignment = al("center")
for col, val in [(6, round(sum(r[8] or 0 for r in gr_rows), 2)),
                 (7, round(sum(r[9] or 0 for r in gr_rows), 2)),
                 (8, round(gr_total, 2))]:
    c2 = ws6.cell(row=total_row6, column=col, value=val)
    c2.font = ft(bold=True, size=9, color="FFFFFF")
    c2.fill = fill(C["guard_header"])
    c2.number_format = "#,##0.00"
    c2.alignment = al("center")
ws6.row_dimensions[total_row6].height = 18

ws6.freeze_panes = "B5"
ws6.sheet_view.zoomScale = 90

# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 7: 📦 Product Matrix
# ═══════════════════════════════════════════════════════════════════════════════
print("Building Product Matrix...")
ws7 = wb.create_sheet("📦 Product Matrix")
set_col_widths(ws7, [28] + [14]*6 + [16,8,8])

write_title_row(ws7, "  LeKise  |  Product × Month Matrix 2026  |  16 มิ.ย. 2569", 1, "A1:J1")
ws7.row_dimensions[2].height = 6

pm_headers = ["กลุ่มสินค้า / เดือน"] + MONTHS + ["รวม","%","ใบ"]
for col, h in enumerate(pm_headers, start=1):
    c = ws7.cell(row=3, column=col, value=h)
    c.font = ft(bold=True, size=9, color="FFFFFF")
    c.fill = fill(C["header_prod"])
    c.alignment = al("center", wrap=True)
    c.border = border()
ws7.row_dimensions[3].height = 28

for r_i, pg in enumerate(PRODUCTS_ORDER, start=4):
    d = by_product.get(pg, {"count": 0, "total": 0.0})
    row_total = 0
    row_cnt = 0
    bg = C["light_blue"] if r_i % 2 == 0 else C["white"]
    c = ws7.cell(row=r_i, column=1, value=pg)
    c.font = ft(size=9)
    c.fill = fill(bg)
    c.alignment = al("left")
    c.border = border()

    for m_i, m in enumerate(MONTHS, start=2):
        m_total = sum(r[10] or 0 for r in all_data if r[6] == pg and r[1] == m)
        m_cnt   = sum(1 for r in all_data if r[6] == pg and r[1] == m)
        row_total += m_total
        row_cnt += m_cnt
        c2 = ws7.cell(row=r_i, column=m_i, value=round(m_total, 2) if m_total else None)
        c2.font = ft(size=9)
        c2.fill = fill(bg)
        c2.alignment = al("center")
        c2.border = border()
        if m_total:
            c2.number_format = "#,##0"

    pct = row_total / grand_total_all * 100 if grand_total_all else 0
    for col, val, fmt in [(8, round(row_total, 2), "#,##0"), (9, f"{pct:.1f}%", None), (10, row_cnt, None)]:
        c3 = ws7.cell(row=r_i, column=col, value=val)
        c3.font = ft(bold=True, size=9)
        c3.fill = fill(bg)
        c3.alignment = al("center")
        c3.border = border()
        if fmt:
            c3.number_format = fmt
    ws7.row_dimensions[r_i].height = 16

# Month totals row
total_row7 = len(PRODUCTS_ORDER) + 4
ws7.cell(row=total_row7, column=1, value="รวมทุกกลุ่ม")
ws7.cell(row=total_row7, column=1).font = ft(bold=True, size=9, color="FFFFFF")
ws7.cell(row=total_row7, column=1).fill = fill(C["header_prod"])
ws7.cell(row=total_row7, column=1).alignment = al("center")
ws7.cell(row=total_row7, column=1).border = border()

for m_i, m in enumerate(MONTHS, start=2):
    m_total = sum(r[10] or 0 for r in all_data if r[1] == m)
    c2 = ws7.cell(row=total_row7, column=m_i, value=round(m_total, 2))
    c2.font = ft(bold=True, size=9, color="FFFFFF")
    c2.fill = fill(C["header_prod"])
    c2.number_format = "#,##0"
    c2.alignment = al("center")
    c2.border = border()

for col, val, fmt in [(8, round(grand_total_all, 2), "#,##0"), (9, "100%", None), (10, total_records, None)]:
    c3 = ws7.cell(row=total_row7, column=col, value=val)
    c3.font = ft(bold=True, size=9, color="FFFFFF")
    c3.fill = fill(C["header_prod"])
    c3.alignment = al("center")
    c3.border = border()
    if fmt:
        c3.number_format = fmt
ws7.row_dimensions[total_row7].height = 18
ws7.freeze_panes = "B4"
ws7.sheet_view.zoomScale = 90

# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 8: ✅ Follow-Up Tracker
# ═══════════════════════════════════════════════════════════════════════════════
print("Building Follow-Up Tracker...")
ws8 = wb.create_sheet("✅ Follow-Up Tracker")
set_col_widths(ws8, [5,8,26,10,20,18,14,12,16,18,22,14,14,14])

fu_rows = sorted([r for r in all_data if (r[10] or 0) >= 100000],
                 key=lambda x: x[10] or 0, reverse=True)
fu_total = sum(r[10] or 0 for r in fu_rows)

write_title_row(ws8,
    f"  LeKise  |  Follow-Up Tracker  |  Priority Order  |  16 มิ.ย. 2569",
    1, "A1:N1")
ws8.row_dimensions[2].height = 6

ws8.merge_cells("A3:N3")
c = ws8.cell(row=3, column=1,
             value=f"ใบเสนอ ≥ 100,000 บาท  |  {len(fu_rows)} รายการ  |  รวม {fu_total/1e6:.1f}M ฿  |  16 มิ.ย. 2569")
c.font = ft(bold=True, size=9, color="1B5E20")
c.fill = fill(C["light_green"])
c.alignment = al("center")
ws8.row_dimensions[3].height = 18

fu_headers = ["#","PL No.","ชื่อลูกค้า","เดือน","สินค้าหลัก","เลขที่ใบเสนอ",
              "ก่อน VAT","VAT 7%","รวม incl VAT","สถานะ (Auto)","สถานะตามจริง",
              "ความสำคัญ","กลุ่มงาน","Action"]
header_row(ws8, fu_headers, 4, colors=[C["header_fu"]] * 14)
ws8.row_dimensions[4].height = 28

priority_order = {"🔥": 0, "📞": 1, "📋": 2, "⚪": 3}
fu_rows_sorted = sorted(fu_rows, key=lambda r: (priority_order.get(str(r[11] or "")[:1], 9), -(r[10] or 0)))

for i, r in enumerate(fu_rows_sorted, start=1):
    row_i = i + 4
    status = str(r[11] or "")
    bg = C["row_hot"] if "🔥" in status else C["row_high"] if "📞" in status else C["row_med"] if "📋" in status else C["row_low"]
    action = "📞 ติดตามด่วน!" if "🔥" in status else "📋 นัดประชุม" if "📞" in status else "📧 ส่ง Email" if "📋" in status else "⚪ รอสอบถาม"
    vals = [i, r[0], r[2], r[1], r[7][:30] if r[7] else "", r[4],
            round(r[8] or 0, 2), round(r[9] or 0, 2), round(r[10] or 0, 2),
            r[11], r[15] or "—", r[12], r[14] or "—", action]
    for col, val in enumerate(vals, start=1):
        c = ws8.cell(row=row_i, column=col, value=val)
        c.font = ft(size=9)
        c.fill = fill(bg)
        c.border = border("thin", "CFD8DC")
        c.alignment = al("left" if col in (3, 5, 6) else "center")
        if col in (7, 8, 9):
            c.number_format = "#,##0.00"
    ws8.row_dimensions[row_i].height = 14

# Total
total_row8 = len(fu_rows_sorted) + 5
ws8.merge_cells(f"A{total_row8}:F{total_row8}")
c = ws8.cell(row=total_row8, column=1, value=f"รวม Follow-Up ≥ 100K — {len(fu_rows)} รายการ")
c.font = ft(bold=True, size=9, color="FFFFFF")
c.fill = fill(C["header_fu"])
c.alignment = al("center")
for col, val in [(7, round(sum(r[8] or 0 for r in fu_rows), 2)),
                 (8, round(sum(r[9] or 0 for r in fu_rows), 2)),
                 (9, round(fu_total, 2))]:
    c2 = ws8.cell(row=total_row8, column=col, value=val)
    c2.font = ft(bold=True, size=9, color="FFFFFF")
    c2.fill = fill(C["header_fu"])
    c2.number_format = "#,##0.00"
    c2.alignment = al("center")
ws8.row_dimensions[total_row8].height = 18

ws8.freeze_panes = "C5"
ws8.sheet_view.zoomScale = 90

# ─── Sheet tab colors ─────────────────────────────────────────────────────────
ws1.sheet_properties.tabColor = "1A237E"
ws2.sheet_properties.tabColor = "283593"
ws3.sheet_properties.tabColor = "00695C"
ws4.sheet_properties.tabColor = "4527A0"
ws5.sheet_properties.tabColor = "1565C0"
ws6.sheet_properties.tabColor = "BF360C"
ws7.sheet_properties.tabColor = "37474F"
ws8.sheet_properties.tabColor = "1B5E20"

# ─── Save ─────────────────────────────────────────────────────────────────────
print(f"\nSaving to {OUT} ...")
wb.save(OUT)
print("✅ Done!")

# ─── Summary report ───────────────────────────────────────────────────────────
print("\n" + "="*60)
print("📋 SUMMARY REPORT")
print("="*60)
print(f"  Existing records loaded : 303")
print(f"  Duplicates removed      : 1  (PL857 → replaced by PL940)")
print(f"  New records added       : {len(new_records)}")
print(f"  Total records now       : {total_records}")
print(f"  Grand Pipeline (incl VAT): ฿{grand_total_all:,.2f}")
print(f"  🔥 Hot (≥3M)  : ฿{hot_total:,.2f}")
print(f"  📞 High (1-3M): ฿{high_total:,.2f}")
print(f"  📋 Med (0.3-1M): ฿{follow_total:,.2f}")
print(f"  ⚪ Low (<0.3M): ฿{wait_total:,.2f}")
print(f"  Unique customers        : {unique_customers}")
print("\n  New PL added:")
for r in new_records:
    print(f"    {r[0]}: {r[2][:30]:<30} | {r[6]:<22} | ฿{r[10]:>12,.2f} | {r[11]}")
print(f"\n  Removed (duplicate):")
for r in removed_info:
    print(f"    {r[0]}: {str(r[2]):<30} | {str(r[6]):<22} | ฿{r[10]:>12,.2f}")
print("="*60)
